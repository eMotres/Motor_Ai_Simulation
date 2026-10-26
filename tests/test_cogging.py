"""Cogging torque as a configuration-level parameter (owner 2026-10-26).

Pins: the pk-pk / half-amplitude / % of continuous arithmetic, the cogging order (LCM), the kwargs the no-load run asks the
solver for (I = 0, cogging-quality sampling, virtual-work torque, eddy off), saving on the configuration + staleness, and
every place the number is shown (tree, project comparison keys, references, datasheet, report rows).  No FEM: the solver is
replaced by a function that returns a known waveform.
"""
from __future__ import annotations

import io
import math
from pathlib import Path

import pytest
import yaml
from fastapi.testclient import TestClient

from motor_ai_sim import cogging as C
from motor_ai_sim.api import app

_ROOT = Path(__file__).resolve().parents[1]
_REAL_DIES = _ROOT / "config" / "dies"
_REAL_USERS = _ROOT / "config" / "users.json"

DIE = "TESTDIE COG"
CFG = "L40"
ADMIN = "admin@example.com"

client = TestClient(app)


# ── pure numbers ─────────────────────────────────────────────────────────────

def _wave(n=144, amp=0.02, cycles=12, offset=0.0):
    return [offset + amp * math.sin(2 * math.pi * cycles * i / n) for i in range(n)]


def test_cogging_order_is_lcm_and_tiles_the_electrical_period():
    c = C.cogging_cycles(12, 14)
    assert (c["order"], c["pole_pairs"], c["cycles_per_electrical_period"]) == (84, 7, 12)
    assert C.cogging_cycles(12, 10)["order"] == 60 and C.cogging_cycles(12, 10)["cycles_per_electrical_period"] == 12
    assert C.cogging_cycles(9, 6)["cycles_per_electrical_period"] == 6
    assert C.cogging_cycles(24, 28)["order"] == 168
    assert C.steps_for(12, 14) == 144 and C.steps_for(12, 14, 6) == 72
    with pytest.raises(ValueError):
        C.cogging_cycles(12, 13)


def test_slots_poles_from_totals_or_segments():
    assert C.slots_poles({"num_slots": 12, "num_poles": 14}) == (12, 14)
    assert C.slots_poles({"num_seg": 2, "num_slots_per_segment": 6, "num_poles_per_segment": 7}) == (12, 14)
    with pytest.raises(ValueError):
        C.slots_poles({})


def test_summary_pkpk_half_amplitude_fundamental_and_percent():
    s = C.summarize(_wave(), num_slots=12, num_poles=14, continuous_torque_nm=4.0)
    assert s["pkpk_Nm"] == pytest.approx(0.04, rel=1e-3)
    assert s["half_amp_Nm"] == pytest.approx(0.02, rel=1e-3)
    assert s["mean_Nm"] == pytest.approx(0.0, abs=1e-12)
    assert s["fundamental_order"] == 84 and s["cycles_per_electrical_period"] == 12
    assert s["fundamental_amp_Nm"] == pytest.approx(0.02, rel=1e-6)       # a pure fundamental: the DFT line equals the amplitude
    assert s["dominant_order"] == 84
    assert s["samples_per_cycle"] == 12.0
    assert s["pct_of_continuous"] == pytest.approx(1.0, rel=1e-3)         # 0.04 / 4.0
    assert "pct_of_continuous" not in C.summarize(_wave(), num_slots=12, num_poles=14)
    assert "pct_of_continuous" not in C.summarize(_wave(), num_slots=12, num_poles=14, continuous_torque_nm=0.0)


def test_summary_refuses_a_series_it_cannot_judge():
    with pytest.raises(ValueError):
        C.summarize([0.0, 1.0], num_slots=12, num_poles=14)
    with pytest.raises(ValueError):
        C.summarize([0.0, float("nan"), 1.0, 2.0], num_slots=12, num_poles=14)


# ── the run asks the solver for the right thing ──────────────────────────────

def _fake_solver(seen):
    def fake(**kw):
        seen.append(kw)
        n = kw["n_steps_per_period"]
        per = 360.0 / 7
        return {"T_em_Nm": _wave(n, 0.03, n // 12),
                "rotor_angle_deg": [per * i / n for i in range(n)],
                "torque_method": "coulomb_virtual_work", "frozen_nu": kw.get("frozen_nu"),
                "gap_layers_effective": 3.0}
    return fake


def _die_doc():
    return {"name": DIE, "locked": False, "created": "2026-10-26T10:00:00",
            "geometry": {"num_slots": 12, "num_poles": 14, "stator_diameter": 40.0,
                         "magnet_height": 3.0, "motor_length": 40.0}}


def _cfg_doc(**extra):
    d = {"name": CFG, "die": DIE, "role": "motor",
         "geometry_overrides": {"motor_length": 40.0, "wire_height": 1.0},
         "winding": {"connection": "2S"},
         "materials": {"magnet": "F52SH_120C", "stator_core": "20SW1200"},
         "duties": [{"name": "rated", "mode": "motor", "saved_at": "2026-10-26T10:00:00",
                     "current_arms": 85.0, "rpm": 6000.0, "gamma_deg": 12.0, "torque_nm": 4.0,
                     "summary": {"T_em_avg_Nm": 4.0}, "note": ""}]}
    d.update(extra)
    return d


def test_run_asks_for_a_no_load_virtual_work_cogging_solve():
    seen = []
    machine = C.machine_from_config(_die_doc(), _cfg_doc(), {"mesh_size_mm": 2.0})
    run = C.run(machine, continuous_torque_nm=4.0, em_transient_eval=_fake_solver(seen))
    kw = seen[0]
    assert kw["I_phase_rms"] == 0.0 and kw["eddy"] is False and kw["rotor_eddy"] is False and kw["demag"] is False
    assert kw["sampling_purpose"] == "cogging_quality" and kw["n_periods"] == 1.0
    assert kw["n_steps_per_period"] == 144                                  # 12 samples x 12 cycles per electrical period
    assert kw["torque_method"] == "coulomb" and kw["element_order"] == 2 and kw["drive"] == "current"
    assert kw["daxis_deg"] == 0.0                                           # given: no calibration solve is paid for
    assert kw["frozen_nu"] is False                                         # the iron is re-converged (see the module note)
    assert kw["mesh_size_mm"] == 2.0 and kw["geo_override"]["num_slots"] == 12
    assert run["summary"]["pkpk_Nm"] == pytest.approx(0.06, rel=1e-3)
    assert run["summary"]["pct_of_continuous"] == pytest.approx(1.5, rel=1e-3)
    assert run["magnet_card"] == "F52SH_120C" and run["torque_method"] == "coulomb_virtual_work"
    assert run["geo_hash"] and run["waveform"]["torque_Nm"] and "re-converged" in run["method"]
    # the frozen-permeability variant stays available for comparison and says so
    run2 = C.run(machine, frozen_nu=True, em_transient_eval=_fake_solver(seen))
    assert seen[1]["frozen_nu"] is True and run2["frozen_nu"] is True and "frozen permeability" in run2["method"]
    with pytest.raises(ValueError):
        C.run(machine, samples_per_cycle=3, em_transient_eval=_fake_solver(seen))


def test_sector_count_is_the_largest_divisor_of_gcd_not_above_the_request():
    assert C._clamp_sectors(4, 12, 14) == 2          # gcd 2
    assert C._clamp_sectors(4, 24, 28) == 4
    assert C._clamp_sectors(8, 24, 28) == 4
    assert C._clamp_sectors(4, 9, 8) == -1           # gcd 1: the full machine


# ── the saved block, its view and staleness ──────────────────────────────────

def _run_for(die_doc, cfg_doc, pk=0.03):
    machine = C.machine_from_config(die_doc, cfg_doc)
    r = C.run(machine, em_transient_eval=_fake_solver([]))
    r["build_sig"] = C.current_build_sig(die_doc, cfg_doc)
    return r


def test_view_is_flat_small_and_hides_nothing_it_should_not():
    dd, cd = _die_doc(), _cfg_doc()
    run = _run_for(dd, cd)
    blk = C.build_block(run, build_sig=C.current_build_sig(dd, cd))
    v = C.view(blk, current_sig=C.current_build_sig(dd, cd), cfg_doc=cd)
    assert v["pkpk_Nm"] == pytest.approx(0.06, rel=1e-3) and v["half_amp_Nm"] == pytest.approx(0.03, rel=1e-3)
    assert v["stale"] is False and v["pct_of_continuous"] == pytest.approx(1.5, rel=1e-3)
    assert v["continuous_torque_Nm"] == 4.0 and "solved torque of duty 'rated'" in v["continuous_torque_source"]
    assert "waveform" not in v
    flat = C.flat_keys(v)
    assert flat["cogging_pkpk_nm"] == v["pkpk_Nm"] and flat["cogging_stale"] is False and flat["cogging_order"] == 84
    # nothing / zero / not a number is "nothing": no empty or zero rows anywhere
    assert C.view(None) is None and C.flat_keys(None) == {}
    assert C.view({"pkpk_Nm": 0}) is None and C.view({"pkpk_Nm": "x"}) is None and C.view({"pkpk_Nm": float("nan")}) is None
    # a block with no stamp, or judged against no signature, is stale (unknown is never "fine")
    assert C.view({"pkpk_Nm": 0.01}, current_sig="x")["stale"] is True
    assert C.view(dict(blk), current_sig=None)["stale"] is True


def test_continuous_torque_prefers_the_s1_search_then_the_rated_duty():
    cd = _cfg_doc()
    stored = {"rated": {"coupled": {"continuous_rating": {"ok": True, "feasible": True, "power": {"T_em_Nm": 3.2}}}}}
    assert C.continuous_torque(cd, stored) == (3.2, "continuous rating (S1 search) of duty 'rated'")
    assert C.continuous_torque(cd, {})[0] == 4.0
    assert C.continuous_torque({"duties": []}, {}) == (None, None)


def test_build_signature_moves_with_the_build_and_marks_the_block_stale():
    dd, cd = _die_doc(), _cfg_doc()
    sig = C.current_build_sig(dd, cd)
    blk = C.build_block(_run_for(dd, cd), build_sig=sig)
    assert not C.is_stale(blk, sig)
    cd2 = _cfg_doc(materials={"magnet": "N52UH_150C", "stator_core": "20SW1200"})      # the magnet card changed
    assert C.is_stale(blk, C.current_build_sig(dd, cd2))
    cd3 = _cfg_doc(geometry_overrides={"motor_length": 41.0, "wire_height": 1.0})
    assert C.is_stale(blk, C.current_build_sig(dd, cd3))
    # a duty edit is not a build change
    cd4 = _cfg_doc()
    cd4["duties"][0]["rpm"] = 7000.0
    assert not C.is_stale(blk, C.current_build_sig(dd, cd4))


# ── the routes on a throwaway catalog ────────────────────────────────────────

@pytest.fixture(scope="module", autouse=True)
def _real_catalog_untouched():
    before = {p: p.stat().st_mtime_ns for p in sorted(_REAL_DIES.rglob("*")) if p.is_file()}
    yield
    after = {p: p.stat().st_mtime_ns for p in sorted(_REAL_DIES.rglob("*")) if p.is_file()}
    assert before == after, "config/dies was modified by a test"


@pytest.fixture()
def dies(tmp_path, monkeypatch):
    from motor_ai_sim.routes import family as fam
    root = tmp_path / "dies"
    (root / DIE).mkdir(parents=True)
    (root / DIE / "die.yaml").write_text(yaml.safe_dump(_die_doc(), sort_keys=False, allow_unicode=True), encoding="utf-8")
    (root / DIE / f"{CFG}.yaml").write_text(yaml.safe_dump(_cfg_doc(), sort_keys=False, allow_unicode=True), encoding="utf-8")
    monkeypatch.setattr(fam, "_DIES_DIR", root)
    return root


@pytest.fixture()
def granted(tmp_path, monkeypatch, dies):
    import shutil
    from motor_ai_sim import auth
    from motor_ai_sim import users as U
    users_file = tmp_path / "users.json"
    if _REAL_USERS.exists():
        shutil.copy2(_REAL_USERS, users_file)
    monkeypatch.setattr(U, "_USERS_FILE", users_file)
    monkeypatch.setenv("AUTH_SECRET", "test-secret-not-the-real-one")
    monkeypatch.delenv("CATALOG_GRANT_ALL_REGISTERED", raising=False)
    monkeypatch.setattr(auth, "_ADMIN_EMAILS", {ADMIN})
    monkeypatch.setattr(auth, "AUTH_ENFORCE", False)
    U.create_user(ADMIN, "password-admin", role="admin", name="Admin")
    return {"Authorization": f"Bearer {U.issue_token(ADMIN)}"}


def _cfg_on_disk(dies):
    return yaml.safe_load((dies / DIE / f"{CFG}.yaml").read_text(encoding="utf-8"))


@pytest.fixture()
def fake_solver(monkeypatch):
    from motor_ai_sim import solve_pool
    seen = []
    monkeypatch.setattr(solve_pool, "em_transient_eval", _fake_solver(seen))
    return seen


def test_run_with_save_files_the_block_and_reads_back(dies, granted, fake_solver):
    # nothing saved yet: the read says so and draws nothing
    r = client.get(f"/api/simulation/cogging/{DIE}/{CFG}", headers=granted)
    assert r.status_code == 200 and r.json()["cogging"] is None
    r = client.post("/api/simulation/cogging/run", headers=granted,
                    json={"die": DIE, "config": CFG, "save": True, "mesh_size_mm": 2.0})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["saved"] is True and out["cogging_id"].startswith("cog_")
    assert out["summary"]["pkpk_Nm"] == pytest.approx(0.06, rel=1e-3)
    assert fake_solver[0]["I_phase_rms"] == 0.0 and fake_solver[0]["mesh_size_mm"] == 2.0
    blk = _cfg_on_disk(dies)["cogging"]
    assert blk["pkpk_Nm"] == pytest.approx(0.06, abs=1e-4) and blk["fundamental_order"] == 84
    assert blk["build_sig"] and blk["magnet_card"] == "F52SH_120C" and len(blk["waveform"]["torque_Nm"]) == 144
    got = client.get(f"/api/simulation/cogging/{DIE}/{CFG}?waveform=true", headers=granted).json()
    assert got["cogging"]["stale"] is False and got["cogging"]["pct_of_continuous"] == pytest.approx(1.5, rel=1e-2)
    assert len(got["waveform"]["torque_Nm"]) == 144
    # the family tree carries the view (the Motors row)
    tree = client.get("/api/family/tree", headers=granted).json()
    cfg_row = next(c for d in tree["dies"] if d["name"] == DIE for c in d["configs"] if c["name"] == CFG)
    assert cfg_row["cogging"]["pkpk_Nm"] == pytest.approx(0.06, abs=1e-4) and cfg_row["cogging"]["stale"] is False


def test_the_saved_value_goes_stale_when_the_build_changes(dies, granted, fake_solver):
    assert client.post("/api/simulation/cogging/run", headers=granted,
                       json={"die": DIE, "config": CFG, "save": True}).status_code == 200
    doc = _cfg_on_disk(dies)
    doc["materials"]["magnet"] = "N52UH_150C"                               # the magnet card changes
    (dies / DIE / f"{CFG}.yaml").write_text(yaml.safe_dump(doc, sort_keys=False, allow_unicode=True), encoding="utf-8")
    got = client.get(f"/api/simulation/cogging/{DIE}/{CFG}", headers=granted).json()
    assert got["cogging"]["stale"] is True                                  # still shown, marked superseded
    tree = client.get("/api/family/tree", headers=granted).json()
    cfg_row = next(c for d in tree["dies"] if d["name"] == DIE for c in d["configs"] if c["name"] == CFG)
    assert cfg_row["cogging"]["stale"] is True


def test_a_run_without_save_changes_nothing_and_can_be_saved_by_id(dies, granted, fake_solver):
    before = (dies / DIE / f"{CFG}.yaml").read_bytes()
    out = client.post("/api/simulation/cogging/run", headers=granted, json={"die": DIE, "config": CFG}).json()
    assert out["saved"] is False and (dies / DIE / f"{CFG}.yaml").read_bytes() == before
    r = client.put(f"/api/simulation/cogging/{DIE}/{CFG}", headers=granted, json={"cogging_id": out["cogging_id"]})
    assert r.status_code == 200 and r.json()["saved"] is True
    assert _cfg_on_disk(dies)["cogging"]["pkpk_Nm"] == pytest.approx(0.06, abs=1e-4)
    assert client.put(f"/api/simulation/cogging/{DIE}/{CFG}", headers=granted,
                      json={"cogging_id": "cog_nope"}).status_code == 404


def test_a_save_is_refused_when_the_configuration_changed_after_the_run(dies, granted, fake_solver):
    out = client.post("/api/simulation/cogging/run", headers=granted, json={"die": DIE, "config": CFG}).json()
    doc = _cfg_on_disk(dies)
    doc["geometry_overrides"]["motor_length"] = 55.0
    (dies / DIE / f"{CFG}.yaml").write_text(yaml.safe_dump(doc, sort_keys=False, allow_unicode=True), encoding="utf-8")
    r = client.put(f"/api/simulation/cogging/{DIE}/{CFG}", headers=granted, json={"cogging_id": out["cogging_id"]})
    assert r.status_code == 409
    assert "cogging" not in _cfg_on_disk(dies)


def test_save_needs_a_configuration_and_a_half_named_one_is_refused(dies, granted, fake_solver):
    assert client.post("/api/simulation/cogging/run", headers=granted, json={"save": True}).status_code == 422
    assert client.post("/api/simulation/cogging/run", headers=granted, json={"die": DIE}).status_code == 422


def test_the_project_motor_rows_carry_the_flat_cogging_keys(dies, granted, fake_solver):
    from motor_ai_sim.development import motors
    assert client.post("/api/simulation/cogging/run", headers=granted,
                       json={"die": DIE, "config": CFG, "save": True}).status_code == 200
    rows = motors._results(DIE, CFG, _cfg_on_disk(dies), {})
    cmp_ = rows[0]["comparison"]
    assert cmp_["cogging_pkpk_nm"] == pytest.approx(0.06, abs=1e-4)
    assert cmp_["cogging_half_nm"] == pytest.approx(0.03, abs=1e-4)
    assert cmp_["cogging_stale"] is False and cmp_["cogging_order"] == 84
    assert cmp_["cogging_pct_of_continuous"] == pytest.approx(1.5, rel=1e-2)
    # no value saved -> no keys at all (never a 0)
    bare = motors._results(DIE, CFG, _cfg_doc(), {})
    assert not any(k.startswith("cogging_") for k in bare[0]["comparison"])


def test_the_references_answer_carries_the_view_of_a_full_card(dies, granted, fake_solver, monkeypatch):
    from motor_ai_sim.routes import catalog
    assert catalog._cogging_of_card(None) is None and catalog._cogging_of_card({"die": DIE}) is None
    assert catalog._cogging_of_card({"die": DIE, "config": CFG}) is None            # nothing saved yet
    assert client.post("/api/simulation/cogging/run", headers=granted,
                       json={"die": DIE, "config": CFG, "save": True}).status_code == 200
    v = catalog._cogging_of_card({"die": DIE, "config": CFG})
    assert v["pkpk_Nm"] == pytest.approx(0.06, abs=1e-4) and v["stale"] is False


def test_the_restamp_after_a_materials_adoption_keeps_the_block_current():
    from motor_ai_sim.routes import family
    c = _cfg_doc(cogging={"pkpk_Nm": 0.03, "build_sig": "old"})
    assert family._restamp_results(c, "old", "new") == 1 and c["cogging"]["build_sig"] == "new"


# ── the screens that print it ────────────────────────────────────────────────

def _view(stale=False):
    dd, cd = _die_doc(), _cfg_doc()
    blk = C.build_block(_run_for(dd, cd), build_sig=C.current_build_sig(dd, cd))
    return C.view(blk, current_sig=("other" if stale else C.current_build_sig(dd, cd)), cfg_doc=cd)


def test_datasheet_row_sits_beside_the_ripple_and_is_absent_without_a_value():
    from openpyxl import load_workbook
    from motor_ai_sim.datasheet import build_datasheet
    dd = {"geometry": {"stator_diameter": 40, "motor_length": 40, "num_slots": 12, "num_poles": 14}}

    def rows(cfg_doc):
        ws = load_workbook(io.BytesIO(build_datasheet(die="D", cfg="C", die_doc=dd, cfg_doc=cfg_doc)))["Motor card"]
        return [(str(r[0].value), [c.value for c in r[1:]]) for r in ws.iter_rows() if r[0].value]

    base = _cfg_doc()
    assert not any("Cogging" in k for k, _ in rows(base))
    blk = C.build_block(_run_for(dd, base), build_sig=C.current_build_sig(dd, base))
    labels = [k for k, _ in rows(_cfg_doc(cogging=blk))]
    assert "Cogging torque, peak-to-peak (N·m)" in labels
    assert labels.index("Cogging torque, peak-to-peak (N·m)") == labels.index("Torque ripple (%)") + 1
    row = dict(rows(_cfg_doc(cogging=blk)))["Cogging torque, peak-to-peak (N·m)"]
    assert row[0] == pytest.approx(0.06, abs=1e-4)
    note = next(v for v in row if isinstance(v, str))
    assert "half-amplitude 0.03" in note and "% of the continuous torque" in note and "older build" not in note
    # a build that moved since says so in the note
    blk_old = dict(blk, build_sig="not-the-current-build")
    note2 = next(v for v in dict(rows(_cfg_doc(cogging=blk_old)))["Cogging torque, peak-to-peak (N·m)"] if isinstance(v, str))
    assert "older build" in note2


def test_report_rows_print_cogging_in_both_tables_and_only_when_saved():
    from motor_ai_sim import report as R
    v = _view()
    em = {"T_em_avg_Nm": 4.0, "T_ripple_pct": 2.0, "rpm": 6000.0}
    labels = [r[0] for r in R.em_torque_rows(em)]
    assert "Cogging torque, peak-to-peak [N·m]" not in labels
    tp = R.em_torque_rows(em, cogging=v)
    labels = [r[0] for r in tp]
    assert labels.index("Cogging torque, peak-to-peak [N·m]") == labels.index("Torque ripple [%]") + 1
    row = next(r for r in tp if r[0] == "Cogging torque, peak-to-peak [N·m]")
    assert row[1] == "0.06" and "half-amplitude 0.03" in row[2] and "no current" in row[2]
    col = {"duty": "rated", "d": {"rpm": 6000.0, "mode": "motor"}, "em": dict(em), "result": {}, "res": {}, "cogging": v}
    _h, rows = R.em_compare_rows([col], {})
    assert dict((r[0], r[1]) for r in rows)["Cogging torque, peak-to-peak [N·m]"] == "0.06"
    col2 = dict(col, cogging=None)
    assert "Cogging torque, peak-to-peak [N·m]" not in [r[0] for r in R.em_compare_rows([col2], {})[1]]
    assert "older build" in R.cogging_words(_view(stale=True))


def test_the_headless_script_runs_in_process_and_saves(dies, fake_solver, capsys):
    """scripts/cogging_run.py - the call a server queue job / a container uses (no login, no HTTP)."""
    import argparse
    import importlib.util
    spec = importlib.util.spec_from_file_location("cogging_run", _ROOT / "scripts" / "cogging_run.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    a = argparse.Namespace(die=DIE, config=CFG, save=True, mesh_size=2.0, magnet_temp=None, frozen_nu=False,
                           samples_per_cycle=12)
    out = mod._local(a)
    mod._print(out)
    text = capsys.readouterr().out
    assert "cogging torque (peak-to-peak): 0.06 N·m" in text and "SAVED on the configuration" in text
    doc = yaml.safe_load((dies / DIE / f"{CFG}.yaml").read_text(encoding="utf-8"))
    assert doc["cogging"]["pkpk_Nm"] == pytest.approx(0.06, abs=1e-4)
    # a second call without --save leaves the file alone
    before = (dies / DIE / f"{CFG}.yaml").read_bytes()
    a.save = False
    assert mod._local(a)["saved"] is False and (dies / DIE / f"{CFG}.yaml").read_bytes() == before
