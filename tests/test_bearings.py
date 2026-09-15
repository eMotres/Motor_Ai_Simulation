"""Bearing friction and rotor windage — the model, the library, and the route.

The point of this file is the FIRST test in it.  ``motor_ai_sim.bearings`` is not
a new correlation: it is the SKF frictional-moment model, and that model was
already pinned against a bench measurement on this project's own hardware —
``docs/measurements/2026-08-04_noload_decomposition_150mm.md`` §4.1, the real
150 mm 24s/28p on 2 x SKF 61811-2RS1, where the measured free-run curve returned
the bearing torque as 0.387 +/- 0.025 N.m against the model's 0.400 N.m with
nothing fitted.  If ``test_150mm_validated_pair`` goes red, the number the
machine was measured against has moved and the loss picture is wrong.

The rest is the machinery around it: the loader following the file (the failure
mode ``materials._load`` exists to prevent), the speed check, windage's power
laws, and the route answering honestly when a machine has no bearings.
"""
from __future__ import annotations

import math

import pytest

from motor_ai_sim import bearings as brg


# ---------------------------------------------------------------------------
# 1.  THE VALIDATED CASE — 2 x SKF 61811-2RS1 on the real 150 mm
# ---------------------------------------------------------------------------
# Doc §4.1, every constant pinned from the SKF tables:
#   seals   K_S1 = 0.018, K_S2 = 20, beta = 2.25, d_s = 59.5 mm -> 197.0 N.mm
#   rolling R1 = 4.7e-7, G_rr = R1.d_m^1.96.F_r^0.54, M_rr = G_rr.(n.nu)^0.6
#           with d_m = 63.5 mm, nu = 30 mm^2/s, F_r = 6.51 N
#   sliding S1 = 6.50e-3, G_sl = S1.d_m^-0.26.F_r^(5/3), mu_sl = 0.05
# The document computed with Phi_ish = Phi_rs = 1 (they are within a few percent
# of unity at 2000 rpm on this bearing); this module applies the full SKF
# factors, which is why the tolerance is +/-0.01 N.m rather than exact.
NU_VALIDATED = 30.0          # mm^2/s, grease base oil at the running temperature
F_R_VALIDATED = 6.51         # N per bearing — 1.328 kg rotating, split over two


def _pair_torque(rpm: float) -> float:
    card = brg.get_bearing("61811-2RS1")
    one = brg.friction(card, rpm=rpm, f_r_n=F_R_VALIDATED,
                       nu_mm2_s=NU_VALIDATED)
    return 2.0 * float(one["M_total_Nm"])


def test_150mm_validated_pair_torque():
    """The bench number: 0.400 N.m at 2000 rpm, 0.404 at 4000."""
    assert _pair_torque(2000.0) == pytest.approx(0.400, abs=0.01)
    assert _pair_torque(4000.0) == pytest.approx(0.404, abs=0.01)


def test_150mm_seal_term_is_the_whole_story():
    """M_seal = 197.0 N.mm per bearing, and it is ~98 % of the torque.

    That is WHY P_bearing is very nearly linear in speed on this machine, which
    is the fact the free-run decomposition leans on to separate it from the n^2
    term it was hunting.
    """
    card = brg.get_bearing("61811-2RS1")
    f = brg.friction(card, rpm=2000.0, f_r_n=F_R_VALIDATED, nu_mm2_s=NU_VALIDATED)
    assert f["M_seal_Nm"] * 1e3 == pytest.approx(197.0, abs=0.5)
    assert f["M_seal_Nm"] / f["M_total_Nm"] > 0.97
    # the rolling term, per the doc: 0.0032 N.m at 2000 rpm before the factors
    assert f["M_rr_Nm"] == pytest.approx(0.0032, abs=0.0006)
    # the sliding term is negligible and must stay so — 2.5e-6 N.m
    assert f["M_sl_Nm"] == pytest.approx(2.5e-6, rel=0.2)


def test_150mm_bearing_power_at_4000rpm():
    """169 W at 4000 rpm (doc §6) — the number that made the whole exercise
    worth doing, since the computed iron loss at the same speed is 80 W."""
    p = _pair_torque(4000.0) * 4000.0 * 2 * math.pi / 60.0
    assert p == pytest.approx(169.2, rel=0.03)


def test_load_barely_matters():
    """M_rr goes as F_r^0.54 and the seal term not at all, so tripling the
    rotating mass moves a sealed bearing's torque by ~1 % (doc §4.1 says
    "< 1 %"; the model gives 1.2 %, because the doc rounded 197.0 + 3.25 N.mm
    against 197.0 + 5.88 — the CONCLUSION, that an unweighed rotor is good
    enough here, is the same either way).  This is what lets the loss picture
    be trusted without a scale under the rotor."""
    a = brg.friction(brg.get_bearing("61811-2RS1"), rpm=2000.0,
                     f_r_n=F_R_VALIDATED, nu_mm2_s=NU_VALIDATED)["M_total_Nm"]
    b = brg.friction(brg.get_bearing("61811-2RS1"), rpm=2000.0,
                     f_r_n=3 * F_R_VALIDATED, nu_mm2_s=NU_VALIDATED)["M_total_Nm"]
    assert abs(b - a) / a < 0.02


# ---------------------------------------------------------------------------
# 2.  Seals: the difference between a shield and a lip
# ---------------------------------------------------------------------------

def test_shields_have_no_seal_drag():
    """2Z metal shields do not touch the inner ring, so M_seal is EXACTLY zero —
    not a small number.  The 40 mm motor's 618/8-2Z is the opposite regime to
    the sealed cards: milliwatts of rolling loss and nothing else."""
    f = brg.friction(brg.get_bearing("618/8-2Z"), rpm=20000.0, f_r_n=2.0,
                     temp_c=60.0)
    assert f["M_seal_Nm"] == 0.0
    assert f["M_total_Nm"] < 1e-3          # well under a milli-newton-metre
    assert not f["seal_estimate"]


def test_low_friction_seal_is_flagged_as_an_estimate():
    """2RZ gets half the RS1 value and SAYS it is an estimate — a number the
    reader must be able to discount."""
    f = brg.friction(brg.get_bearing("6010-2RZ/HC5C3"), rpm=3000.0, f_r_n=50.0,
                     temp_c=70.0)
    assert f["seal_estimate"] is True
    assert f["M_seal_Nm"] > 0.0
    assert any("ESTIMATE" in n for n in f["notes"])


# ---------------------------------------------------------------------------
# 3.  Viscosity
# ---------------------------------------------------------------------------

def test_walther_reproduces_its_own_anchor_points():
    nu = brg.walther_nu(110.0, 11.0, 40.0)
    assert nu == pytest.approx(110.0, rel=1e-6)
    assert brg.walther_nu(110.0, 11.0, 100.0) == pytest.approx(11.0, rel=1e-6)
    # and it falls monotonically in between, which is the whole point
    assert 11.0 < brg.walther_nu(110.0, 11.0, 70.0) < 110.0


def test_walther_refuses_a_swapped_pair():
    """A grease card with nu100 above nu40 is a typo, and interpolating it would
    make the bearing get STIFFER as it warms up."""
    with pytest.raises(ValueError):
        brg.walther_nu(11.0, 110.0, 70.0)


def test_the_validated_grease_gives_nu_30_at_a_plausible_temperature():
    """MT47 is what SKF fills a 2RS1 with, and the free-run decomposition ran at
    nu = 30 mm^2/s.  The catalogue card and the measurement have to be
    reconcilable, and they are: ~59 C, an ordinary bearing temperature."""
    lube = brg.get_lubricant("MT47")
    assert lube.nu_at(59.0) == pytest.approx(30.0, abs=1.0)


# ---------------------------------------------------------------------------
# 4.  The loader follows the file
# ---------------------------------------------------------------------------

def test_loader_reloads_on_mtime_and_survives_a_broken_edit(tmp_path, monkeypatch):
    """The failure ``materials._load`` was written to stop: a card edited on disk
    stayed invisible until a restart, so a run silently used the wrong part.
    Same contract here — and a half-written file keeps the last good copy rather
    than taking the process down mid-request."""
    import time as _t

    lib = tmp_path / "bearings_library.yaml"
    lib.write_text(
        "bearings:\n"
        "  TEST-1:\n"
        "    type: deep_groove\n"
        "    d: 10\n    D: 30\n    B: 9\n"
        "    seals: none\n"
        "    default_lubricant: L\n"
        "    friction: {R1: 4.3e-7, S1: 4.25e-3}\n"
        "lubricants:\n"
        "  L: {nu40: 100, nu100: 10}\n", encoding="utf-8")
    monkeypatch.setattr(brg, "_LIB_PATH", lib, raising=True)
    monkeypatch.setattr(brg, "_library", None, raising=True)
    monkeypatch.setattr(brg, "_lib_mtime", 0.0, raising=True)
    monkeypatch.setattr(brg, "_lib_checked", 0.0, raising=True)
    brg._CARD_CACHE.clear()
    brg._LUBE_CACHE.clear()

    assert brg.get_bearing("TEST-1").D == 30

    # (a) an EDIT is picked up — after the probe window, and only then
    _t.sleep(1.1)
    lib.write_text(
        "bearings:\n"
        "  TEST-1:\n"
        "    type: deep_groove\n"
        "    d: 10\n    D: 47\n    B: 9\n"
        "    seals: none\n"
        "    default_lubricant: L\n"
        "    friction: {R1: 4.3e-7, S1: 4.25e-3}\n"
        "lubricants:\n"
        "  L: {nu40: 100, nu100: 10}\n", encoding="utf-8")
    assert brg.get_bearing("TEST-1").D == 47, "the cache did not follow the file"

    # (b) a MALFORMED edit keeps the last good copy
    _t.sleep(1.1)
    lib.write_text("bearings:\n  TEST-1:\n   d: [unclosed\n", encoding="utf-8")
    assert brg.get_bearing("TEST-1").D == 47

    # (c) an unknown name names what IS there rather than falling back
    with pytest.raises(brg.UnknownBearingError) as e:
        brg.get_bearing("NOT-A-BEARING")
    assert "TEST-1" in str(e.value)


def _pin_library(monkeypatch, path):
    monkeypatch.setattr(brg, "_LIB_PATH", path, raising=True)
    monkeypatch.setattr(brg, "_library", None, raising=True)
    monkeypatch.setattr(brg, "_lib_mtime", 0.0, raising=True)
    monkeypatch.setattr(brg, "_lib_checked", 0.0, raising=True)
    brg._CARD_CACHE.clear()
    brg._LUBE_CACHE.clear()


_CARD_TMPL = ("bearings:\n"
              "  T-1:\n"
              "    d: 10\n    D: 30\n    B: 9\n    seals: none\n"
              "    stiffness_n_per_m: {stiff}\n"
              "    default_lubricant: L\n"
              "    friction: {{R1: 4.3e-7, S1: 4.25e-3}}\n"
              "lubricants:\n  L: {{nu40: 100, nu100: 10}}\n")


def test_a_yaml11_unsigned_exponent_still_lands_as_a_number(tmp_path, monkeypatch):
    """YAML 1.1 parses ``1.0e8`` as the STRING "1.0e8" — a float literal needs a
    signed exponent.  A stiffness that is secretly a string reaches the
    rotordynamics default as text and fails three layers from the typo, so the
    card loader coerces it.  (Not hypothetical: the shipped library was written
    that way, and this is how it was found.)"""
    lib = tmp_path / "bearings_library.yaml"
    lib.write_text(_CARD_TMPL.format(stiff="1.0e8"), encoding="utf-8")
    _pin_library(monkeypatch, lib)
    k = brg.get_bearing("T-1").stiffness_n_per_m
    assert isinstance(k, float) and k == pytest.approx(1.0e8)


def test_a_non_numeric_catalogue_value_is_refused_by_name(tmp_path, monkeypatch):
    """Text that is not a number at all cannot be repaired, and the refusal
    names the card AND the field — an engineer-readable message, per the
    project's standing validation rule."""
    lib = tmp_path / "bearings_library.yaml"
    lib.write_text(_CARD_TMPL.format(stiff='"about 1e8 N/m"'), encoding="utf-8")
    _pin_library(monkeypatch, lib)
    with pytest.raises(ValueError) as e:
        brg.get_bearing("T-1")
    assert "T-1" in str(e.value) and "stiffness_n_per_m" in str(e.value)


def test_the_shipped_library_parses_and_every_card_computes():
    """Every catalogue card must survive a friction call: a card with a missing
    constant is a machine whose losses can never be computed, and it should be
    found here, not by a user."""
    names = brg.list_bearings()
    assert "61811-2RS1" in names and "71910 CE/HCP4A" in names
    for n in names:
        card = brg.get_bearing(n)
        f = brg.friction(card, rpm=1000.0, f_r_n=30.0, f_a_n=50.0, temp_c=60.0)
        assert f["M_total_Nm"] > 0.0
        assert f["P_W"] > 0.0
        assert card.stiffness_n_per_m and card.stiffness_n_per_m > 0


# ---------------------------------------------------------------------------
# 5.  Speed check
# ---------------------------------------------------------------------------

def test_speed_check_grease_vs_oil_air_on_the_200mm_candidate():
    """The 71910 at 23 000 rpm is over its grease limit and inside its oil-air
    one.  That single fact is the reason the Ø200 machine cannot be greased."""
    card = brg.get_bearing("71910 CE/HCP4A")
    g = brg.speed_check(card, 23000.0, "grease")
    o = brg.speed_check(card, 23000.0, "oil_air")
    assert g["ok"] is False and "OVER" in g["verdict"]
    assert o["ok"] is True
    # n.dm is the number the industry compares bearings with
    assert g["n_dm"] == pytest.approx(23000.0 * 61.0)


def test_speed_check_says_so_when_a_card_has_no_limit_for_that_lubricant():
    """A sealed, grease-for-life bearing has no oil-air limit, and answering
    'ok' would be inventing permission."""
    s = brg.speed_check(brg.get_bearing("61811-2RS1"), 3000.0, "oil_air")
    assert s["ok"] is None and s["limit_rpm"] is None
    assert "no oil-air limit" in s["verdict"]


def test_the_sealed_hybrid_is_the_counter_example():
    """6010-2RZ/HC5C3 has the same envelope and the same ceramic balls as the
    7010 CE/HCP4A, and still cannot go near 23 000 rpm — the SEAL caps it."""
    sealed = brg.speed_check(brg.get_bearing("6010-2RZ/HC5C3"), 23000.0, "grease")
    ac = brg.speed_check(brg.get_bearing("7010 CE/HCP4A"), 23000.0, "oil_air")
    assert sealed["ok"] is False
    assert ac["ok"] is True


# ---------------------------------------------------------------------------
# 6.  Windage
# ---------------------------------------------------------------------------

_W150 = dict(r_rotor_m=0.0563, r_bore_m=0.0568, length_m=0.035, temp_c=50.0)


def test_windage_matches_the_validated_150mm_figure():
    """1.12 W at 4000 rpm in the decomposition, from a different gap model
    (laminar concentric cylinder + Taylor enhancement).  The two must agree to
    within the model spread, or one of them is wrong."""
    w = brg.windage(rpm=4000.0, **_W150)
    assert w["P_W"] == pytest.approx(1.12, rel=0.20)
    # and it is SMALL: under 0.5 % of the 555 W budget at that speed
    assert w["P_W"] < 3.0


def _exponent(f, n1: float, n2: float) -> float:
    return math.log(f(n2) / f(n1)) / math.log(n2 / n1)


def test_windage_laminar_gap_torque_is_linear_and_power_quadratic():
    """In the laminar Couette regime C_f = 2/Re, so C_f.omega^2 collapses to
    omega: TORQUE goes as n^1 and POWER as n^2.  (This is the classic
    M = 2.pi.mu.omega.r^3.L/delta, and getting it wrong by an order in omega is
    the easiest possible mistake in a drag model.)"""
    def gap_torque(n):
        w = brg.windage(rpm=n, **_W150)
        assert "laminar" in w["gap_regime"], w["gap_regime"]
        return w["M_gap_Nm"]

    def gap_power(n):
        return brg.windage(rpm=n, **_W150)["P_gap_W"]

    assert _exponent(gap_torque, 200.0, 600.0) == pytest.approx(1.0, abs=0.05)
    assert _exponent(gap_power, 200.0, 600.0) == pytest.approx(2.0, abs=0.05)


def test_windage_turbulent_gap_torque_exponent():
    """Above the Taylor threshold the Wendt/Bilgen-Boulos coefficients put the
    gap torque between n^1.5 and n^1.7 — no longer viscous, not yet n^2."""
    def gap_torque(n):
        w = brg.windage(rpm=n, **_W150)
        assert "laminar (Ta" not in w["gap_regime"], w["gap_regime"]
        return w["M_gap_Nm"]

    k = _exponent(gap_torque, 3000.0, 6000.0)
    assert 1.4 < k < 1.8, f"gap torque exponent {k}"


def test_windage_end_faces_follow_the_disc_law():
    """Laminar disc C_M = 3.87/Re^0.5 -> torque n^1.5; turbulent
    C_M = 0.146/Re^0.2 -> n^1.8."""
    big = dict(r_rotor_m=0.0657, r_bore_m=0.0663, length_m=0.180, temp_c=70.0)

    def faces(n, g):
        return brg.windage(rpm=n, **g)["M_faces_Nm"]

    assert _exponent(lambda n: faces(n, _W150), 500.0, 1500.0) \
        == pytest.approx(1.5, abs=0.05)
    # the 200 mm rotor at 20-30 krpm is well past Re = 3e5
    w = brg.windage(rpm=23000.0, **big)
    assert w["face_regime"] == "turbulent disc"
    assert _exponent(lambda n: faces(n, big), 20000.0, 30000.0) \
        == pytest.approx(1.8, abs=0.05)


def test_windage_at_rest_is_zero_and_finite():
    w = brg.windage(rpm=0.0, **_W150)
    assert w["P_W"] == 0.0 and w["M_total_Nm"] == 0.0


def test_windage_from_geometry_refuses_an_incomplete_geometry():
    assert brg.windage_from_geometry({"rotor_outer_radius": 63.2}, 1000.0) is None
    g = {"rotor_outer_radius": 63.2, "air_gap": 3.1, "motor_length": 180,
         "sleeve_thickness": 2.5}
    w = brg.windage_from_geometry(g, 1000.0)
    assert w is not None
    # the clearance is air_gap - sleeve = 0.6 mm, NOT the 3.1 mm air gap
    assert w["delta_mm"] == pytest.approx(0.6, abs=1e-6)


# ---------------------------------------------------------------------------
# 7.  The machine-level roll-up
# ---------------------------------------------------------------------------

_ASSIGN_150 = {
    "A": {"card": "61811-2RS1"},
    "B": {"card": "61811-2RS1"},
    "lubrication": "grease",
    "preload_n": 0,
    "temp_source": "manual",
    "temp_c": 59,
}

_GEO_150 = {"rotor_outer_radius": 56.3, "air_gap": 0.5, "motor_length": 35,
            "sleeve_thickness": 0}


def test_machine_rollup_reproduces_the_pair_and_names_what_it_omits():
    out = brg.machine_bearing_losses(_ASSIGN_150, rpm=2000.0, temp_c=59.0,
                                     rotor_mass_kg=1.328, geometry=_GEO_150)
    assert out["has_bearings"] is True
    assert len(out["bearings"]) == 2
    assert out["M_bearings_Nm"] == pytest.approx(0.400, abs=0.015)
    assert out["P_bearings_W"] == pytest.approx(84.0, rel=0.05)
    assert out["P_windage_W"] < 1.0
    assert out["P_mech_extra_W"] == pytest.approx(
        out["P_bearings_W"] + out["P_windage_W"])
    # the omission has to be STATED, every time
    assert any("Unbalanced magnetic pull" in n for n in out["notes"])


def test_machine_rollup_splits_the_weight_over_the_span():
    """With a span and an offset the reactions follow the lever rule and still
    sum to the rotor's weight."""
    beam = {"bearing_span_mm": 200, "stack_offset_mm": 50}
    out = brg.machine_bearing_losses(_ASSIGN_150, rpm=2000.0, temp_c=59.0,
                                     rotor_mass_kg=10.0, geometry=_GEO_150,
                                     beam=beam)
    fa = out["bearings"][0]["F_r_N"]
    fb = out["bearings"][1]["F_r_N"]
    assert fa + fb == pytest.approx(10.0 * 9.80665, rel=1e-9)
    assert fb > fa                       # the stack moved toward bearing B
    assert any("lever rule" in n for n in out["notes"])


def test_machine_rollup_without_bearings_is_not_zero_loss():
    out = brg.machine_bearing_losses({}, rpm=2000.0, rotor_mass_kg=1.3)
    assert out["has_bearings"] is False
    assert "P_mech_extra_W" not in out           # NOT a computed zero
    assert "Shaft & bearings" in out["note"]


def test_one_card_named_means_two_of_it():
    out = brg.machine_bearing_losses({"A": {"card": "61811-2RS1"}}, rpm=2000.0,
                                     temp_c=59.0, rotor_mass_kg=1.328)
    assert len(out["bearings"]) == 2
    assert out["bearings"][1]["bearing"] == "61811-2RS1"


def test_200mm_candidate_is_over_its_grease_limit_and_says_so():
    """The engineering answer for the hollow-shaft Ø200 at 23 000 rpm: the
    71910 pair works on oil-air and not on grease, and the roll-up must carry
    that verdict into its notes rather than only into a sub-object."""
    a = {"A": {"card": "71910 CE/HCP4A"}, "B": {"card": "71910 CE/HCP4A"},
         "lubrication": "grease", "preload_n": 200}
    out = brg.machine_bearing_losses(a, rpm=23000.0, temp_c=70.0,
                                     rotor_mass_kg=10.2)
    assert any("OVER the limit" in n for n in out["notes"])
    # the centrifugal ball load is the dominant term up here, and it is reported
    assert out["bearings"][0]["F_g_N"] > out["bearings"][0]["F_r_N"] * 10


def test_hybrid_ball_factor_lowers_the_centrifugal_term():
    """F_g is a ball-MASS term, so a Si3N4 ball makes it 0.41x a steel one. The
    factor is applied openly and reported — a hybrid that came out identical to
    a steel bearing would mean the field is being ignored."""
    f = brg.friction(brg.get_bearing("71910 CE/HCP4A"), rpm=20000.0,
                     f_r_n=50.0, f_a_n=200.0, temp_c=70.0)
    assert f["ball_density_ratio"] == pytest.approx(0.408)
    raw = 1.90e-12 * 61.0 ** 4 * 20000.0 ** 2
    assert f["F_g_N"] == pytest.approx(raw * 0.408, rel=1e-6)


def test_a_deep_groove_card_never_swallows_an_axial_load_silently():
    """The 618-series cards carry only SKF's radial row.  Given a preload they
    must carry it AND say what that costs in accuracy — a dropped load would be
    an efficiency that is quietly too good."""
    f = brg.friction(brg.get_bearing("61814-2RS1"), rpm=3000.0, f_r_n=20.0,
                     f_a_n=300.0, temp_c=70.0)
    assert any("equivalent radial" in n for n in f["notes"])
    bare = brg.friction(brg.get_bearing("61814-2RS1"), rpm=3000.0, f_r_n=20.0,
                        temp_c=70.0)
    assert f["M_total_Nm"] > bare["M_total_Nm"]


# ---------------------------------------------------------------------------
# 8.  The routes
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def client():
    from fastapi.testclient import TestClient

    from motor_ai_sim.api import app
    return TestClient(app)


def test_library_route_serves_every_card(client):
    r = client.get("/api/bearings/library")
    assert r.status_code == 200
    j = r.json()
    assert "61811-2RS1" in j["bearings"]
    card = j["bearings"]["61811-2RS1"]
    assert card["d"] == 55 and card["D"] == 72 and card["d_m"] == 63.5
    assert card["friction"]["R1"] == pytest.approx(4.7e-7)
    assert "MT47" in j["lubricants"]


def test_losses_route_with_an_in_memory_assignment(client, monkeypatch):
    """The route reads the MACHINE's block; here it is monkeypatched in so the
    test never touches a die file (which this suite must not write)."""
    from motor_ai_sim.routes import bearings as route

    monkeypatch.setattr(
        route, "_machine_bearings",
        lambda die, cfg: (dict(_ASSIGN_150), "TESTDIE", "L35"), raising=True)
    monkeypatch.setattr(route, "_live_geometry", lambda: dict(_GEO_150),
                        raising=True)

    r = client.get("/api/bearings/losses",
                   params={"rpm": 2000, "temp_c": 59, "rotor_mass_kg": 1.328})
    assert r.status_code == 200
    j = r.json()
    assert j["has_bearings"] is True
    assert j["die"] == "TESTDIE" and j["config"] == "L35"
    assert j["P_bearings_W"] == pytest.approx(84.0, rel=0.06)
    assert j["M_bearings_Nm"] == pytest.approx(0.400, abs=0.015)
    assert j["windage"]["P_W"] < 1.0
    assert j["rotor_mass_source"] == "given by the caller"


def test_losses_route_says_no_bearings_rather_than_zero(client, monkeypatch):
    from motor_ai_sim.routes import bearings as route

    monkeypatch.setattr(route, "_machine_bearings",
                        lambda die, cfg: (None, "D", "C"), raising=True)
    monkeypatch.setattr(route, "_live_geometry", lambda: dict(_GEO_150),
                        raising=True)
    j = client.get("/api/bearings/losses", params={"rpm": 2000}).json()
    assert j["has_bearings"] is False
    assert "P_bearings_W" not in j
    # windage is a property of the ROTOR, so it is still answered
    assert j["P_windage_W"] is not None
    assert "Shaft & bearings" in j["note"]


def test_losses_route_refuses_nonsense(client):
    assert client.get("/api/bearings/losses", params={"rpm": 0}).status_code == 422
    assert client.get("/api/bearings/losses",
                      params={"rpm": 1000, "lubrication": "petrol"}).status_code == 422


def test_the_patch_route_refuses_before_it_writes(client):
    """A die file naming a bearing nobody has is a machine whose losses can
    never be computed, and the 422 would arrive weeks later — so the card is
    checked at the SAVE.  It is checked BEFORE the yaml is opened, which is also
    why this test cannot write anything: the refusal comes first.
    (The same order guards the lubrication and preload validation.)"""
    for body, says in (
        ({"A": {"card": "NOT-A-BEARING"}}, "unknown bearing 'NOT-A-BEARING'"),
        ({"A": {"card": "61811-2RS1"}, "lubrication": "petrol"}, "'grease' or 'oil_air'"),
        ({"A": {"card": "61811-2RS1"}, "preload_n": -5}, "preload_n must be non-negative"),
        ({"A": {"card": "61811-2RS1"}, "temp_c": 9000}, "outside anything a bearing survives"),
        ({"A": {"card": "61811-2RS1"}, "temp_source": "vibes"}, "'manual' or 'thermal'"),
    ):
        r = client.patch("/api/family/config/nope-die/nope-cfg/bearings", json=body)
        assert r.status_code == 422, f"{says}: got {r.status_code} {r.text[:200]}"
        # …and the message is one an engineer can act on, not "invalid input"
        assert says in str(r.json().get("detail")), r.text[:300]


def test_losses_route_rejects_an_unknown_card(client, monkeypatch):
    from motor_ai_sim.routes import bearings as route

    monkeypatch.setattr(
        route, "_machine_bearings",
        lambda die, cfg: ({"A": {"card": "NOPE-1"}}, "D", "C"), raising=True)
    r = client.get("/api/bearings/losses", params={"rpm": 1000})
    assert r.status_code == 422
    assert "NOPE-1" in str(r.json()["detail"])


# ---------------------------------------------------------------------------
# 9.  The datasheet stops saying "not included"
# ---------------------------------------------------------------------------

def test_datasheet_prints_mechanical_rows_only_with_bearings():
    """The whole point of the feature, checked where the user reads it."""
    from openpyxl import load_workbook
    import io

    from motor_ai_sim.datasheet import build_datasheet

    die_doc = {"geometry": {"stator_diameter": 150, "motor_length": 35,
                            "num_slots": 24, "num_poles": 28,
                            "rotor_outer_radius": 56.3, "air_gap": 0.5}}
    duty = {"name": "cont", "mode": "motor", "rpm": 2000, "torque_nm": 5.0,
            "result": {"efficiency_pct": 90.0, "loss_w": 120.0},
            "summary": {"P_mech_W": 1047.2, "mass_components": [
                {"name": "Rotor back-iron (X)", "mass_kg": 0.538},
                {"name": "Magnets (Y)", "mass_kg": 0.722},
                {"name": "Shaft (Al)", "mass_kg": 0.067}]}}

    def _labels(cfg_doc) -> set:
        wb = load_workbook(io.BytesIO(build_datasheet(
            die="D", cfg="C", die_doc=die_doc, cfg_doc=cfg_doc)))
        ws = wb["Motor card"]
        return {str(c.value) for c in ws["A"] if c.value}

    bare = _labels({"duties": [duty]})
    assert "Bearing loss (W)" not in bare
    assert "Shaft efficiency (%)" not in bare

    withb = _labels({"duties": [duty], "bearings": _ASSIGN_150})
    assert {"Bearing loss (W)", "Windage (W)",
            "Total losses incl. mechanical (W)",
            "Shaft efficiency (%)"} <= withb
    assert "Efficiency (%)" in withb          # the EM row stays, untouched


# ---------------------------------------------------------------------------
# 10.  ONE implementation — motor_ai_sim.mech_losses (2026-09-08)
# ---------------------------------------------------------------------------
# User: "when the coupled run runs, the WHOLE model must be solved, and ALL the
# losses must be carried into the electromagnetic calculation".  Until this date
# the bearing friction and the rotor windage existed in ONE place — a React
# component fetching /api/bearings/losses — and the stored run, the thermal
# solve and the coupled loop knew nothing about them.  Now four consumers share
# one module, and these tests are what stops them drifting apart again.

def test_the_shared_module_answers_exactly_what_the_route_answers():
    """The route's arithmetic and the summary's arithmetic are ONE function.

    A pair that costs 84 W on the Electromagnetic tab must cost 84 W in the
    thermal map and 84 W in the report, or the difference is a bug in one place
    instead of a disagreement between four.
    """
    from motor_ai_sim import mech_losses as ml

    direct = brg.machine_bearing_losses(
        _ASSIGN_150, rpm=2000.0, temp_c=59.0, rotor_mass_kg=1.328,
        geometry=_GEO_150)
    shared = ml.machine_mech_losses(
        rpm=2000.0, assignment=_ASSIGN_150, geometry=_GEO_150,
        rotor_mass_kg_=1.328, resolve_machine=False)
    assert shared is not None
    for k in ("P_bearings_W", "P_windage_W", "P_mech_extra_W", "M_bearings_Nm"):
        assert shared[k] == pytest.approx(direct[k], rel=1e-12), k
    # …and it says WHICH temperature it used and where that came from, which the
    # bare model does not have to.
    assert shared["bearing_temp_c"] == pytest.approx(59.0)
    assert shared["bearing_temp_source"] == "assigned"


def test_a_machine_with_no_bearings_is_none_not_a_zero():
    """The house rule, at the level every consumer reads.

    ``None`` is what makes the summary fields ABSENT rather than 0 W — an
    unknown mechanical loss printed as zero is an efficiency nobody measured.
    """
    from motor_ai_sim import mech_losses as ml

    assert ml.machine_mech_losses(rpm=2000.0, assignment={},
                                  resolve_machine=False) is None
    assert ml.machine_mech_losses(rpm=0.0, assignment=_ASSIGN_150,
                                  resolve_machine=False) is None
    assert ml.summary_block(None) == {}


def test_the_bearing_temperature_names_its_own_source():
    """Four sources, in the order a reader expects, and never a silent default.

    M_rr goes as nu^0.6, so a grease quoted at 40 degC running at 90 degC is a
    factor of two on the rolling term — a number that important may not be
    invented quietly, which is why the last case is labelled ``default`` and
    says so.
    """
    from motor_ai_sim import mech_losses as ml

    t, src, _note = ml.resolve_bearing_temp(_ASSIGN_150, override_c=93.4)
    assert (t, src) == (93.4, "coupled")

    t, src, _note = ml.resolve_bearing_temp(_ASSIGN_150)
    assert (t, src) == (59.0, "assigned")

    t, src, note = ml.resolve_bearing_temp({"A": {"card": "61811-2RS1"}})
    assert src == "default" and t == 70.0
    assert "70" in note


def test_the_thermal_map_is_read_the_same_way_wherever_it_comes_from():
    """The shaft ENDS first, the shaft average second, nothing third.

    The seat is the metal the exposed-stub conductance acts on, so when that
    path was on its own mean is the measurement.  With it off the shaft is
    adiabatic along the axis and its average is the best the cross-section can
    say.  A map with no shaft at all returns ``None`` — the caller falls back to
    the assignment rather than inventing a seat temperature.
    """
    from motor_ai_sim import mech_losses as ml

    both = {"cooling": {"shaft_ends": {"t_shaft_mean_c": 88.0}},
            "components": {"shaft": {"avg": 61.0, "max": 70.0}}}
    assert ml.bearing_temp_from_map(both)[0] == pytest.approx(88.0)
    assert "exposed ends" in ml.bearing_temp_from_map(both)[1]

    closed = {"cooling": {"shaft_ends": {"mode": "off",
                                         "t_shaft_mean_c": None}},
              "components": {"shaft": {"avg": 61.0}}}
    assert ml.bearing_temp_from_map(closed)[0] == pytest.approx(61.0)
    assert "cross-section" in ml.bearing_temp_from_map(closed)[1]

    assert ml.bearing_temp_from_map(
        {"components": {"winding": {"avg": 90.0}}}) is None
    assert ml.bearing_temp_from_map(None) is None


def test_thermal_source_uses_the_map_only_when_it_is_this_machine(monkeypatch):
    """A temperature borrowed from another motor is WORSE than the assignment's
    own number, because it looks computed."""
    from motor_ai_sim import mech_losses as ml
    from motor_ai_sim.routes import thermal as th

    a = dict(_ASSIGN_150, temp_source="thermal")
    entry = {"geometry_fingerprint": "MINE",
             "result": {"components": {"shaft": {"avg": 104.0}}}}
    monkeypatch.setattr(th, "_load_last", lambda: None, raising=True)
    monkeypatch.setattr(th, "_LAST", {"field": entry}, raising=True)

    t, src, _ = ml.resolve_bearing_temp(a, geometry_fingerprint="MINE")
    assert (round(t, 1), src) == (104.0, "thermal")
    # another machine's map: the assignment wins back
    t, src, _ = ml.resolve_bearing_temp(a, geometry_fingerprint="SOMEBODY-ELSE")
    assert (t, src) == (59.0, "assigned")


# ---------------------------------------------------------------------------
# 11.  …and the RUN carries them (routes.simulation._mech_loss_fields)
# ---------------------------------------------------------------------------

_SBRES = {"geo_fingerprint": "FP-150"}
_MASS_ROWS = [{"name": "Rotor back-iron (steel)", "mass_kg": 0.538},
              {"name": "Magnets (N52UH)", "mass_kg": 0.722},
              {"name": "Shaft (Al)", "mass_kg": 0.068}]


@pytest.fixture
def machine_with_bearings(monkeypatch):
    """The ACTIVE machine, in memory: the 150 mm pair and its geometry.

    Monkeypatched at ``mech_losses`` rather than at the route, so the code under
    test is the real resolution path and only the die-file read is faked — this
    suite may not touch ``config/dies``.
    """
    from motor_ai_sim import mech_losses as ml

    monkeypatch.setattr(ml, "machine_bearings",
                        lambda die=None, cfg=None: (dict(_ASSIGN_150),
                                                    "TESTDIE", "L35"),
                        raising=True)
    monkeypatch.setattr(ml, "live_geometry",
                        lambda ov=None: dict(_GEO_150), raising=True)
    return ml


def _fields(**over):
    from motor_ai_sim.routes import simulation as sim

    kw = dict(sbres=dict(_SBRES), geo_override=None, rpm=2000.0,
              p_mech_w=1047.2, p_loss_w=120.0, efficiency=0.897,
              op_mode="motor", mass_components=list(_MASS_ROWS))
    kw.update(over)
    sbres = kw.pop("sbres")
    return sim._mech_loss_fields(sbres, **kw)


def test_a_stored_run_of_this_machine_carries_the_mechanical_half(
        machine_with_bearings):
    """The user's requirement, at the field level.

    Every one of these used to be computed in the BROWSER and thrown away; now
    they ride with the run, so the datasheet, the report, Compare and the
    coupled loop read the same numbers the card shows.
    """
    f = _fields()
    assert f["P_bearings_W"] == pytest.approx(84.0, rel=0.06)
    assert f["P_windage_W"] < 1.0
    assert f["P_mech_extra_W"] == pytest.approx(
        f["P_bearings_W"] + f["P_windage_W"], abs=0.02)
    assert f["bearing_temp_c"] == pytest.approx(59.0)
    assert f["bearing_temp_source"] == "assigned"
    # what a calorimeter around the ASSEMBLED machine would read
    assert f["P_loss_total_incl_mech_W"] == pytest.approx(
        120.0 + f["P_mech_extra_W"], abs=0.1)
    # motoring: the friction comes OFF the shaft
    assert f["P_shaft_net_W"] == pytest.approx(1047.2 - f["P_mech_extra_W"],
                                               abs=0.1)
    x = f["P_mech_extra_W"] / 1047.2
    assert f["efficiency_shaft"] == pytest.approx(0.897 * (1.0 - x), abs=5e-4)
    assert f["efficiency_shaft"] < 0.897
    # …and it says WHICH model, WHICH cards and at WHAT speed
    m = f["mech_losses"]
    assert m["rpm"] == pytest.approx(2000.0)
    assert m["cards"] == ["61811-2RS1", "61811-2RS1"]
    assert "SKF" in m["model"] and "ANALYTIC" in m["model"]
    assert len(m["bearings"]) == 2 and m["bearings"][0]["end"] == "A"


def test_the_generator_convention_puts_the_friction_on_the_input(
        machine_with_bearings):
    """Generating, the coupling must SUPPLY the rotor's power plus the friction
    that never reaches it — so the shaft sees more than the rotor does, and the
    shaft efficiency is eta/(1+x) rather than eta*(1-x)."""
    f = _fields(op_mode="generator")
    x = f["P_mech_extra_W"] / 1047.2
    assert f["P_shaft_net_W"] == pytest.approx(1047.2 + f["P_mech_extra_W"],
                                               abs=0.1)
    assert f["efficiency_shaft"] == pytest.approx(0.897 / (1.0 + x), abs=5e-4)
    assert "PLUS" in f["P_shaft_convention"]


def test_a_machine_without_bearings_grows_no_fields_at_all(monkeypatch):
    """ABSENT, not zero — the same rule the /losses route follows.  A summary
    with ``P_bearings_W: 0`` would make every consumer print a bearing loss
    nobody measured."""
    from motor_ai_sim import mech_losses as ml

    monkeypatch.setattr(ml, "machine_bearings",
                        lambda die=None, cfg=None: (None, "D", "C"),
                        raising=True)
    assert _fields() == {}


def test_an_optimizer_candidate_is_never_billed_for_bearings(
        machine_with_bearings):
    """Two gates, and both are the point.

    A BACKGROUND solve is a search probe, not a machine anybody has chosen
    bearings for — and paying a die read plus a CAD mass measurement per
    candidate is exactly the cost those gates exist to avoid.  A CANDIDATE
    GEOMETRY is the same statement about a per-request override: its fingerprint
    says it is not the motor whose die file the bearings are written in.
    """
    from motor_ai_sim.routes import simulation as sim

    tok = sim._BACKGROUND_RUN.set(True)
    try:
        assert _fields() == {}
    finally:
        sim._BACKGROUND_RUN.reset(tok)

    # a per-request geometry whose fingerprint is not the live machine's
    assert _fields(geo_override={"stator_diameter": 999.0},
                   sbres={"geo_fingerprint": "SOMETHING-ELSE"}) == {}


def test_a_stored_run_reads_back_in_the_shape_the_report_expects(
        machine_with_bearings):
    """``from_summary`` is what lets the report and the datasheet prefer the
    stored numbers over a recomputation — same shape as the SKF roll-up, plus a
    flag saying it was not recomputed."""
    from motor_ai_sim import mech_losses as ml

    f = _fields()
    back = ml.from_summary(f)
    assert back["has_bearings"] is True and back["from_stored_run"] is True
    assert back["P_bearings_W"] == f["P_bearings_W"]
    assert back["temp_c"] == pytest.approx(59.0)
    assert back["rpm"] == pytest.approx(2000.0)
    # the report's bearing table reads these through `speed`
    sp = back["bearings"][0]["speed"]
    assert set(sp) == {"ok", "verdict", "n_dm", "limit_rpm"}
    assert back["windage"]["M_total_Nm"] is not None
    # a summary from before this change has nothing to read
    assert ml.from_summary({"P_loss_total_W": 120.0}) is None
    assert ml.from_summary(None) is None
