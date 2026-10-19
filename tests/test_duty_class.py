"""SHORT-TIME OR CONTINUOUS — owner 2026-10-19: «помечай как кратковременный».

CIANO28 150_35 new / L35 "rated" reaches its 200 °C winding 6 min 11 s from
cold and has no S1 rating at its speed.  What is pinned here:

* the class is read off the stored coupled record (``duty_class.classify``)
  and named the same way everywhere ("rated — short-time: 6 min 11 s from
  cold (winding 200 °C)");
* the compacted coupled record carries it flat;
* a short-time default duty gives way to a continuous one on the catalog row;
* the datasheet leads with the continuous duty and labels the short-time one;
* "Save continuous point as duty" writes the point as a duty, renames the
  short-time source instead of overwriting it, and refuses an unconfirmed
  point.
"""
from __future__ import annotations

import pytest
import yaml
from fastapi.testclient import TestClient

from motor_ai_sim import duty_class as dcl
from motor_ai_sim.api import app
from tests.test_family_duty_runs import (CFG, DIE, _cfg_doc,  # noqa: F401
                                         _real_catalog_untouched, dies,
                                         granted)

client = TestClient(app)

#: The L35 "rated" record's two blocks, as the server stored them on
#: 2026-10-18 20:58 (the numbers the owner saw).
TTL_L35 = {"within_limits": False, "limiting_part": "winding",
           "limits_c": {"winding": 200.0, "magnet": 150.0},
           "starts": {"cold": {"time_to_limit_s": 370.7,
                               "limiting_part": "winding"},
                      "rated": {"time_to_limit_s": None}}}
CR_L35 = {"ok": True, "feasible": False, "limiting_part": "magnet",
          "reference": {"rpm": 4000.0},
          "cooling_label": "propeller tmotor_g40x13_1 slipstream, 40 °C",
          "continuous_point": {"feasible": True, "rpm": 2643.0,
                               "I_cont_A_rms": 24.0, "verified": True,
                               "limiting_part": "magnet", "basis": "load_line",
                               "power": {"P_shaft_W": 2400.0,
                                         "T_em_Nm": 8.7}}}
COUPLED_L35 = {"time_to_limit": TTL_L35, "continuous_rating": CR_L35,
               "computed_at": "2026-10-18T20:58:00"}


def test_the_l35_rated_record_is_short_time_and_says_for_how_long():
    k = dcl.classify({"name": "rated"}, COUPLED_L35)
    assert k["duty_class"] == dcl.SHORT_TIME
    assert k["time_to_limit_s"] == pytest.approx(370.7)
    assert dcl.label("rated", k) == ("rated — short-time: 6 min 11 s from cold "
                                     "(winding 200 °C)")
    assert dcl.fmt_duration(45.2) == "45 s"
    assert dcl.fmt_duration(7500) == "2 h 05 min"


def test_what_makes_a_duty_continuous():
    inside = {"time_to_limit": {"within_limits": True}}
    assert dcl.classify({}, inside)["duty_class"] == dcl.CONTINUOUS
    s1 = {"continuous_rating": {"ok": True, "feasible": True, "s": 1.2,
                                "verified": True}}
    assert dcl.classify({}, s1)["duty_class"] == dcl.CONTINUOUS
    below = {"continuous_rating": {"ok": True, "feasible": True, "s": 0.7}}
    assert dcl.classify({}, below) is None
    saved = {"name": "rated", "duty_class": "continuous"}
    assert dcl.classify(saved, None)["duty_class"] == dcl.CONTINUOUS
    # …and a later run that finds it over its limit wins over the yaml flag.
    assert dcl.classify(saved, COUPLED_L35)["duty_class"] == dcl.SHORT_TIME
    assert dcl.classify({}, None) is None


def test_the_compacted_record_carries_the_class_flat():
    from motor_ai_sim.duty_results import compact_coupled
    rec = compact_coupled({"coupling": dict(COUPLED_L35)})
    assert rec["duty_class"] == dcl.SHORT_TIME
    assert rec["time_to_limit_s"] == pytest.approx(370.7)
    assert rec["short_time_part"] == "winding"
    assert rec["short_time_limit_c"] == 200.0
    # …and reading it back from the flat keys gives the same label.
    assert dcl.label("rated", dcl.classify({}, {
        k: rec[k] for k in ("duty_class", "time_to_limit_s",
                            "short_time_part", "short_time_limit_c")})) \
        == "rated — short-time: 6 min 11 s from cold (winding 200 °C)"


def test_a_short_time_default_gives_way_to_a_continuous_duty():
    from motor_ai_sim.routes.family import config_rated_duty
    c = {"duties": [{"name": "rated"}, {"name": "rated_cont"},
                    {"name": "peak"}]}
    classes = {"rated": {"duty_class": dcl.SHORT_TIME},
               "rated_cont": {"duty_class": dcl.CONTINUOUS},
               "peak": {"duty_class": dcl.SHORT_TIME}}
    assert config_rated_duty(c, classes) == "rated_cont"
    # no continuous duty → the configured default stays (and is labelled)
    assert config_rated_duty(c, {"rated": {"duty_class": dcl.SHORT_TIME}}) \
        == "rated"
    assert config_rated_duty(c, {}) == "rated"


def test_the_datasheet_leads_with_the_continuous_duty(tmp_path):
    from openpyxl import load_workbook

    from motor_ai_sim.datasheet import build_datasheet

    die_doc = {"geometry": {"num_poles": 28, "num_slots": 24,
                            "stator_diameter": 150.0, "motor_length": 35.0}}
    cfg_doc = {"winding": {"star_delta": "star"},
               "duties": [{"name": "rated", "current_arms": 82.3,
                           "rpm": 4000.0, "gamma_deg": 16.0},
                          {"name": "rated_cont", "current_arms": 24.0,
                           "rpm": 2643.0, "gamma_deg": 16.0,
                           "duty_class": "continuous"}]}
    blob = build_datasheet(die="D", cfg="L35", die_doc=die_doc,
                           cfg_doc=cfg_doc,
                           coupled={"rated": dict(COUPLED_L35)})
    p = tmp_path / "card.xlsx"
    p.write_bytes(blob)
    ws = load_workbook(p).active

    def row(label):
        for r in range(1, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value or "") == label:
                return [ws.cell(row=r, column=2).value,
                        ws.cell(row=r, column=3).value]
        return None

    assert row("Duty") == ["rated_cont", "rated (short-time)"]
    assert row("Duty class") == [
        "continuous", "short-time: 6 min 11 s from cold (winding 200 °C)"]


def _stored(monkeypatch, coupled):
    """The duty-results store, as the route reads it — never the real one."""
    from motor_ai_sim import duty_fields, duty_results
    monkeypatch.setattr(duty_results, "get",
                        lambda die, cfg: {"rated": {"coupled": coupled}})
    monkeypatch.setattr(duty_results, "rename", lambda *a, **k: True)
    monkeypatch.setattr(duty_fields, "rename", lambda *a, **k: 0)


def _make_rated(dies):
    p = dies / DIE / f"{CFG}.yaml"
    c = yaml.safe_load(p.read_text(encoding="utf-8"))
    c["duties"] = [{"name": "rated", "mode": "motor",
                    "saved_at": "2026-10-18T20:00:00", "current_arms": 82.3,
                    "rpm": 4000.0, "gamma_deg": 16.0, "note": "",
                    "star_delta": "star", "mesh": {"mesh.maxSize": 2.0}}]
    p.write_text(yaml.safe_dump(c, sort_keys=False, allow_unicode=True),
                 encoding="utf-8")


def test_save_continuous_point_takes_the_name_and_renames_the_source(
        dies, granted, monkeypatch):
    _make_rated(dies)
    _stored(monkeypatch, dict(COUPLED_L35))
    r = client.post(f"/api/family/duty/{DIE}/{CFG}/rated/save_continuous_point",
                    json={"name": "rated"}, headers=granted["admin"])
    assert r.status_code == 200, r.text
    assert r.json()["renamed_source_to"] == "rated_short_time"
    by = {d["name"]: d for d in _cfg_doc(dies)["duties"]}
    assert set(by) == {"rated", "rated_short_time"}
    new = by["rated"]
    assert new["duty_class"] == "continuous"
    assert new["current_arms"] == 24.0 and new["rpm"] == 2643.0
    assert new["gamma_deg"] == 16.0 and new["mesh"] == {"mesh.maxSize": 2.0}
    assert new["continuous_source"]["duty"] == "rated_short_time"
    assert by["rated_short_time"]["current_arms"] == 82.3   # renamed, intact


def test_save_continuous_point_refuses_what_it_cannot_stand_behind(
        dies, granted, monkeypatch):
    _make_rated(dies)
    url = f"/api/family/duty/{DIE}/{CFG}/rated/save_continuous_point"
    # no continuous point at all
    _stored(monkeypatch, {"time_to_limit": TTL_L35})
    assert client.post(url, json={}, headers=granted["admin"]).status_code == 409
    # a point its real pass did not confirm
    unconfirmed = dict(COUPLED_L35, continuous_rating=dict(
        CR_L35, continuous_point=dict(CR_L35["continuous_point"],
                                      verified=False, note="+6.2 K")))
    _stored(monkeypatch, unconfirmed)
    r = client.post(url, json={}, headers=granted["admin"])
    assert r.status_code == 409 and "NOT confirmed" in r.text
    # a reader may not write
    _stored(monkeypatch, dict(COUPLED_L35))
    assert client.post(url, json={}, headers=granted["client"]).status_code \
        in (401, 403, 404)
    assert [d["name"] for d in _cfg_doc(dies)["duties"]] == ["rated"]


def test_the_report_overview_marks_a_short_time_duty():
    from motor_ai_sim.report import duty_class_mark, short_time_mark
    col = {"d": {"name": "rated", "mode": "motor"},
           "res": {"coupled": dict(COUPLED_L35)}}
    assert short_time_mark(col) == " · short-time 6 min 11 s"
    assert duty_class_mark(col) == " · short-time 6 min 11 s"
    assert duty_class_mark({"d": {"name": "rated"}, "res": {}}) == ""
